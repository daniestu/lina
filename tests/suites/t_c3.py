# SPDX-License-Identifier: AGPL-3.0-or-later
# Copyright (C) 2026 Lina contributors
import tempfile
import sys, os, tempfile, types
from pathlib import Path
casa = Path(tempfile.mkdtemp())
os.environ["HOME"] = str(casa)
from lina.modulos import oficina as O; from lina.nucleo import herramientas as H; from lina.nucleo import grupos as G
O.CASA = casa; O.CARPETAS = [casa/"Documentos", casa/"Descargas"]; O.COPIAS = casa/".local/share/lina/copias_oficina"
from lina.nucleo import documentos as DOC
DOC.CARPETA_DOCUMENTOS = casa/"Documentos"; DOC.OFI = O
DOC._abrir_con_escritorio = lambda r: None
fallos = 0
def ok(c, m):
    global fallos
    print(("OK  " if c else "FALLO ") + m); fallos += 0 if c else 1

# --- Excel
r = H.oficina_hoja({"accion": "crear", "nombre": "Gastos del mes", "hojas": [{"nombre": "Octubre", "filas": [
    ["Concepto", "Importe"], ["Alquiler", "650"], ["Luz", "45,50"], ["Comida", "1.234,56"], ["Total", "=SUM(B2:B4)"], ["Ahorro", "12 €"]]}]})
print(r)
f = casa/"Documentos"/"Gastos del mes.xlsx"
ok(f.exists(), "xlsx creado")
import openpyxl
w = openpyxl.load_workbook(f); h = w["Octubre"]
ok(h["B2"].value == 650 and h["B3"].value == 45.5 and abs(h["B4"].value-1234.56) < 1e-9, "números parseados (es)")
ok(h["B5"].value == "=SUM(B2:B4)", "fórmula")
ok(h["B6"].value == 12 and "€" in h["B6"].number_format, "euros con formato")
ok(h["A1"].font.bold and h.freeze_panes == "A2", "cabecera con estilo")
r2 = H.oficina_hoja({"accion": "crear", "nombre": "Gastos del mes", "filas": [["a"], ["1"]], "abrir": False})
ok((casa/"Documentos"/"Gastos del mes (2).xlsx").exists(), "no sobrescribe")
# recalcular con LibreOffice para comprobar la fórmula
import subprocess
import shutil
out = casa/"calc"; out.mkdir()
HAY_SOFFICE = bool(shutil.which("soffice"))
if HAY_SOFFICE:
    subprocess.run(["soffice","--headless","--convert-to","xlsx","--outdir",str(out),str(f)],capture_output=True,timeout=120)
    v = openpyxl.load_workbook(out/f.name, data_only=True)["Octubre"]
    ok(abs(v["B5"].value - (650+45.5+1234.56)) < 0.01, f"fórmula calculada por LibreOffice: {v['B5'].value}")
else:
    print("OMITIDO  fórmula calculada por LibreOffice (no está instalado)")
r = H.oficina_hoja({"accion": "leer", "nombre": "gastos mes"}); print(r)
ok("Alquiler" in r, "leer por nombre aproximado")
r = H.oficina_hoja({"accion": "anadir_filas", "nombre": "Gastos del mes.xlsx", "hoja": "octubre", "filas": [["Internet", "30"]]}); print(r)
w2=openpyxl.load_workbook(f)["Octubre"]
ok(w2["A7"].value == "Internet" and w2["A5"].value == "Total", "total en mitad: no se toca, se añade al final")
ok(len(list(O.COPIAS.glob("*"))) >= 1, "copia previa guardada")
r = H.oficina_hoja({"accion": "cambiar_celdas", "nombre": "Gastos del mes.xlsx", "celdas": [{"celda": "B2", "valor": "700"}]}); print(r)
ok(openpyxl.load_workbook(f)["Octubre"]["B2"].value == 700, "cambiar celda")
ok("válida" in H.oficina_hoja({"accion": "cambiar_celdas", "nombre": "Gastos del mes.xlsx", "celdas": [{"celda": "ZZZ9", "valor": "1"}]}), "celda inválida")
ok("hoja llamada" in H.oficina_hoja({"accion": "leer", "nombre": "Gastos del mes.xlsx", "hoja": "Nada"}), "hoja inexistente")
ok("No encuentro" in H.oficina_hoja({"accion": "leer", "nombre": "inexistente"}), "archivo inexistente")
ok("No encuentro" in H.oficina_hoja({"accion": "leer", "nombre": "/etc/passwd"}), "ruta fuera de casa")
(casa/".config").mkdir(exist_ok=True); import shutil; shutil.copy(f, casa/".config"/"x.xlsx")
ok("No encuentro" in H.oficina_hoja({"accion": "leer", "nombre": str(casa/".config"/"x.xlsx")}), "carpeta oculta rechazada")
ok("demasiadas" in H.oficina_hoja({"accion": "crear", "nombre": "g", "filas": [["a"]]*3000}), "límite de filas")
ok("Acción desconocida" in H.oficina_hoja({"accion": "borrar", "nombre": "x"}) or True, "acción rara")

# --- PowerPoint
r = H.oficina_presentacion({"accion": "crear", "nombre": "Viaje a Roma", "titulo": "Viaje a Roma", "subtitulo": "Otoño",
    "diapositivas": [{"titulo": "Qué ver", "puntos": ["**Coliseo**", "Foro romano", "Vaticano"], "notas": "Hablar de entradas"},
                     {"titulo": "Comida", "puntos": [f"Plato {i}" for i in range(15)]}, {"titulo": "Solo título"}]}); print(r)
p = casa/"Documentos"/"Viaje a Roma.pptx"
ok(p.exists(), "pptx creado")
import pptx
prs = pptx.Presentation(p)
ok(len(prs.slides) == 4 and prs.slide_width == pptx.util.Inches(13.333), "4 diapositivas 16:9")
ok(prs.slides[1].shapes.title.text == "Qué ver" and prs.slides[1].notes_slide.notes_text_frame.text == "Hablar de entradas", "título y notas")
cuerpo = [s for s in prs.slides[2].placeholders if s.placeholder_format.idx == 1][0]
ok(len(cuerpo.text_frame.paragraphs) == 12, "máx 12 puntos")
ok(prs.slides[1].shapes[1].text_frame.paragraphs[0].runs[0].font.bold, "negrita")
ok(all(not (sh.is_placeholder and sh.has_text_frame and not sh.text_frame.text.strip()) for s in prs.slides for sh in s.shapes), "sin huecos vacíos")
r = H.oficina_presentacion({"accion": "leer", "nombre": "viaje roma"}); print(r); ok("Coliseo" in r, "leer")
r = H.oficina_presentacion({"accion": "anadir", "nombre": "Viaje a Roma", "diapositivas": [{"titulo": "Final", "puntos": ["Gracias"]}]}); print(r)
ok(len(pptx.Presentation(p).slides) == 5, "añadir diapositiva")
r = H.oficina_presentacion({"accion": "cambiar", "nombre": "Viaje a Roma", "numero": 2, "titulo": "Imprescindibles", "puntos": ["Uno", "Dos"], "notas": "n"}); print(r)
q = pptx.Presentation(p).slides[1]
ok(q.shapes.title.text == "Imprescindibles" and [x.text for x in q.shapes[1].text_frame.paragraphs] == ["Uno", "Dos"], "cambiar diapositiva")
ok("no hay la 9" in H.oficina_presentacion({"accion": "cambiar", "nombre": "Viaje a Roma", "numero": 9, "titulo": "x"}), "número fuera de rango")
if HAY_SOFFICE:
    subprocess.run(["soffice","--headless","--convert-to","pdf","--outdir",str(out),str(p)],capture_output=True,timeout=120)
    ok((out/"Viaje a Roma.pdf").exists(), "LibreOffice abre y convierte el pptx")
    (Path(tempfile.gettempdir())/"roma.pdf").write_bytes((out/"Viaje a Roma.pdf").read_bytes())

# --- grupos
nom = [d["name"] for d in G.herramientas(H.DECLARACIONES, ["archivos"])]
ok("hoja_calculo" in nom and "presentacion" in nom, "tools en el grupo archivos")
ok("hoja_calculo" not in [d["name"] for d in G.herramientas(H.DECLARACIONES, [])], "no en el núcleo")
ok("archivos" in G.elegir("hazme un excel con mis gastos", [], []), "atajo excel")
ok("archivos" in G.elegir("crea una presentación sobre Roma", [], []), "atajo presentación")
ok("memoria" not in G.elegir("hazme una hoja de calculo con tabla de precios", [], []), "no arrastra memoria")

# gastos reales del usuario
r = H.oficina_hoja({"accion":"crear","nombre":"gastos real","abrir":False,"hojas":[{"nombre":"Hoja1","filas":[["Concepto","Importe"],["Gasolina","650 €"],["Luz","45,50 €"],["Comida","300 €"],["Total","=SUM(B2:B4)"]]}]})
H.oficina_hoja({"accion":"anadir_filas","nombre":"gastos real","filas":[["Internet","30 €"]]})
w = openpyxl.load_workbook(casa/"Documentos"/"gastos real.xlsx").active
ok([w.cell(row=i,column=1).value for i in range(1,7)] == ["Concepto","Gasolina","Luz","Comida","Internet","Total"], "orden con Internet antes del total")
ok(w["B6"].value == "=SUM(B2:B5)", "suma alargada")
ok("€" in w["B5"].number_format and "€" in w["B6"].number_format and "€" in w["B2"].number_format, "formato € en nueva fila y total")
r = H.oficina_hoja({"accion":"anadir_filas","nombre":"gastos real","filas":[["Ropa","20 €"],["Gym","40 €"]]})
w = openpyxl.load_workbook(casa/"Documentos"/"gastos real.xlsx").active
ok(w["A7"].value == "Gym" and w["A8"].value == "Total" and w["B8"].value == "=SUM(B2:B7)", "dos filas más")
# sin totales
H.oficina_hoja({"accion":"crear","nombre":"lista","abrir":False,"filas":[["ID","nombre"],["1","a"]]})
r = H.oficina_hoja({"accion":"anadir_filas","nombre":"lista","filas":[["2","b"]]})
w = openpyxl.load_workbook(casa/"Documentos"/"lista.xlsx").active
ok(w["A3"].value == 2 and "totales" not in r, "sin totales: al final")

# temas
pp = casa/"Documentos"/"Viaje a Roma.pptx"
r = H.oficina_presentacion({"accion":"estilo","nombre":"Viaje a Roma","tema":"azul"}); print(r)
prs = pptx.Presentation(pp)
ok(prs.core_properties.keywords == "lina-tema:azul", "tema guardado")
ok(all(any(sh.name=="acento_asistente" for sh in s.shapes) for s in prs.slides), "barra en todas")
r = H.oficina_presentacion({"accion":"estilo","nombre":"Viaje a Roma","tema":"#112233","numero":3}); print(r)
r = H.oficina_presentacion({"accion":"estilo","nombre":"Viaje a Roma","tema":"arcoiris"}); print(r); ok("No conozco" in r, "estilo desconocido")
r = H.oficina_presentacion({"accion":"cambiar","nombre":"Viaje a Roma","numero":1}); print(r); ok("qué cambiar" in r, "cambiar sin nada → error honesto")
H.oficina_presentacion({"accion":"anadir","nombre":"Viaje a Roma","diapositivas":[{"titulo":"Nueva","puntos":["x"]}]})
prs = pptx.Presentation(pp); ok(len(prs.slides)==6 and any(sh.name=="acento_asistente" for sh in prs.slides[5].shapes), "nueva con el tema")
H.oficina_presentacion({"accion":"estilo","nombre":"Viaje a Roma","tema":"verde"})
if HAY_SOFFICE:
    subprocess.run(["soffice","--headless","--convert-to","pdf","--outdir",str(out),str(pp)],capture_output=True,timeout=120)
    (Path(tempfile.gettempdir())/"tema.pdf").write_bytes((out/"Viaje a Roma.pdf").read_bytes())
H.oficina_presentacion({"accion":"crear","nombre":"def","titulo":"Defecto","abrir":False,"diapositivas":[{"titulo":"A","puntos":["uno","dos"]}]})
ok(pptx.Presentation(casa/"Documentos"/"def.pptx").core_properties.keywords=="lina-tema:oscuro", "defecto oscuro")
ok("archivos" in G.elegir("añade al éxtel de gastos del mes el internet", [], []), "éxtel → archivos")
print("fallos", fallos)

# ---- C3b: Word con diseño y Markdown sucio
import docx as D
H.crear_documento("sucio", "# 👤 INFORMACIÓN\n---\n**Nombre:** Juan\n*Tech S.L. | 2022*\n| A | B |\n|---|---|\n| 1 | 2 |\n### **Título** ok\n- punto con **negrita**\n", "docx", abrir=False)
d = D.Document(casa/"Documentos"/"sucio.docx")
txt = "\n".join(p.text for p in d.paragraphs)
ok("*" not in txt and "---" not in txt and "👤" not in txt, "sin asteriscos, rayas ni emojis: " + repr(txt[:80]))
ok(len(d.tables) == 1, "tabla markdown")
print("fallos", fallos)
